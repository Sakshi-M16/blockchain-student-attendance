"""
generate_documents.py - Generates:
1. YourName_RegNum_BT_Assignment2.docx (Microsoft Word submission document)
2. YourName_RegNum_BT_Assignment2.pdf (PDF submission document)
strictly complying with Assignment 2 submission requirements.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def generate_word_doc(filepath):
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Document Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("Blockchain Technology (CSE4444)")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = subtitle_p.add_run("Assignment 2: Decentralized Application (DApp) Submission")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph() # Spacer

    # Student & Project Details Table
    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    details = [
        ("Student Name:", "[Student Name]"),
        ("Registration Number:", "[Registration Number]"),
        ("Assignment Title:", "Blockchain Technology (CSE4444) - Assignment 2"),
        ("Title of Developed DApp:", "Blockchain-Based Student Attendance Management System"),
        ("GitHub Repository Link:", "https://github.com/[YourUsername]/blockchain-student-attendance")
    ]

    for i, (label, val) in enumerate(details):
        row = table.rows[i]
        
        # Label cell
        cell_lbl = row.cells[0]
        cell_lbl.width = Inches(2.2)
        p_lbl = cell_lbl.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(11)
        r_lbl.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
        set_cell_background(cell_lbl, "F1F5F9")

        # Value cell
        cell_val = row.cells[1]
        cell_val.width = Inches(4.3)
        p_val = cell_val.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(11)
        if "[" in val:
            r_val.font.italic = True
            r_val.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
        set_cell_background(cell_val, "FFFFFF")

    doc.add_paragraph()

    # Section 1: Problem Description
    h1 = doc.add_heading("1. Problem Statement & System Overview", level=1)
    h1.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p1 = doc.add_paragraph(
        "Traditional educational attendance systems rely on centralized databases or paper registers that suffer "
        "from vulnerability to unauthorized modification, proxy attendance, lack of transparency, and risk of data loss. "
        "This project designs and develops an original Decentralized Application (DApp)—the Blockchain-Based Student Attendance "
        "Management System—built on the Ethereum blockchain. By leveraging Ethereum smart contracts, attendance logs are "
        "recorded immutably, ensuring that once attendance is recorded by the authenticated instructor, it cannot be altered "
        "or deleted by any party."
    )
    p1.style.font.size = Pt(11)

    # Section 2: Smart Contract Architecture
    h2 = doc.add_heading("2. Smart Contract Architecture", level=1)
    h2.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p2 = doc.add_paragraph(
        "The smart contract is implemented in Solidity (v0.8.20) under the file contracts/AttendanceSystem.sol. "
        "Key components include:"
    )
    p2.style.font.size = Pt(11)

    bullet_points = [
        ("State Variables & Mappings: ", "Maintains the instructor address, courseName, totalLectures conducted, a Student struct, roll number mappings (mapping(uint256 => Student)), and date-wise attendance mappings (mapping(uint256 => mapping(uint256 => bool)))."),
        ("Constructor: ", "Initializes course name and automatically sets msg.sender as the authorized instructor."),
        ("Validation & Modifiers: ", "Includes the onlyInstructor access control modifier, roll number validation (>0), non-empty string checks, student existence checks, and duplicate attendance prevention."),
        ("Events: ", "Emits StudentRegistered, AttendanceMarked, and LectureConducted events for transparent off-chain indexing."),
        ("Public Views: ", "Exposes getStudent, isAttendedOnDate, getAttendancePercentage, and getAllStudents for student verification.")
    ]

    for title, desc in bullet_points:
        bp = doc.add_paragraph(style='List Bullet')
        r_t = bp.add_run(title)
        r_t.font.bold = True
        r_t.font.size = Pt(10.5)
        r_d = bp.add_run(desc)
        r_d.font.size = Pt(10.5)

    # Section 3: Frontend & MetaMask Integration
    h3 = doc.add_heading("3. User Interface & MetaMask Integration", level=1)
    h3.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p3 = doc.add_paragraph(
        "A lightweight, modern web frontend was created using HTML5, CSS3, and JavaScript, with Ethers.js v5.7.2. "
        "The interface connects directly to MetaMask via the Ethereum provider API (window.ethereum). "
        "It supports both Write Transactions (Instructor registering students and marking attendance with transaction signing) "
        "and Read Calls (Students checking attendance records and attendance percentage without gas fees)."
    )
    p3.style.font.size = Pt(11)

    # Section 4: Testing & Results
    h4 = doc.add_heading("4. Comprehensive Testing & Results", level=1)
    h4.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p4 = doc.add_paragraph(
        "The application was thoroughly tested across valid, invalid, and unauthorized inputs with 100% pass rate:"
    )
    p4.style.font.size = Pt(11)

    # Test Cases Table
    t_test = doc.add_table(rows=1, cols=4)
    t_test.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t_test.rows[0].cells
    headers = ["Test Case", "Input Category", "Scenario & Revert Reason", "Observed Result"]
    for j, h in enumerate(headers):
        hdr[j].text = h
        set_cell_background(hdr[j], "1E3A8A")
        p = hdr[j].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p.runs[0].font.size = Pt(9.5)

    test_data = [
        ("TC-01", "Valid", "Contract Deployment with courseName parameter", "PASS - Deployed, state initialized"),
        ("TC-02", "Valid", "Register Student (Roll 101, 'Alice Johnson')", "PASS - StudentRegistered event emitted"),
        ("TC-03", "Valid", "Mark Attendance (Roll 101, Date 20261017, Present)", "PASS - Attendance recorded, totalAttended=1"),
        ("TC-04", "Valid", "Query Student Attendance & Percentage", "PASS - Correct data displayed on UI"),
        ("TC-05", "Invalid", "Empty Name on Registration: registerStudent(103, '')", "PASS - Revert: 'Student name cannot be empty'"),
        ("TC-06", "Invalid", "Roll number 0: registerStudent(0, 'Charlie')", "PASS - Revert: 'Invalid roll number'"),
        ("TC-07", "Invalid", "Duplicate Student Registration (Roll 101)", "PASS - Revert: 'Student already registered'"),
        ("TC-08", "Invalid", "Mark Attendance for Unregistered Student (Roll 999)", "PASS - Revert: 'Student not registered'"),
        ("TC-09", "Invalid", "Duplicate Attendance on Same Date", "PASS - Revert: 'Attendance already recorded'"),
        ("TC-10", "Unauthorized", "Non-Instructor calls registerStudent", "PASS - Revert: 'Unauthorized: Only instructor'"),
        ("TC-11", "Unauthorized", "Non-Instructor calls markAttendance", "PASS - Revert: 'Unauthorized: Only instructor'")
    ]

    for row_data in test_data:
        r = t_test.add_row().cells
        for col_idx, text in enumerate(row_data):
            r[col_idx].text = text
            p = r[col_idx].paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            if col_idx == 3:
                p.runs[0].font.bold = True
                p.runs[0].font.color.rgb = RGBColor(0x16, 0x65, 0x34)
            set_cell_background(r[col_idx], "F8FAFC" if col_idx % 2 == 0 else "FFFFFF")

    doc.add_paragraph()

    # Section 5: Submission Checklist
    h5 = doc.add_heading("5. Submission Checklist & Repository Contents", level=1)
    h5.style.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    checklist = [
        "contracts/AttendanceSystem.sol - Solidity Smart Contract source file",
        "frontend/ (index.html, style.css, app.js, contract_abi.js) - Complete frontend DApp",
        "deployment/deployment_details.md - Deployment instructions for Remix, Ganache, and Sepolia",
        "test/test_contract.py & test/test_results.md - Automated test suite & test results report",
        "screenshots/ - Organized screenshots folder covering compilation, deployment, MetaMask, and transactions",
        "README.md - Comprehensive documentation with execution & interaction steps"
    ]
    for item in checklist:
        p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(item)
        r.font.size = Pt(10)

    doc.save(filepath)
    print(f"Word document saved to: {filepath}")

def generate_pdf_doc(filepath):
    doc = SimpleDocTemplate(
        filepath,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        alignment=1, # Center
        textColor=colors.HexColor("#1E3A8A")
    )

    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=16,
        alignment=1,
        textColor=colors.HexColor("#475569")
    )

    h1_style = ParagraphStyle(
        'H1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor("#1E3A8A"),
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#0F172A")
    )

    table_cell_style = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11
    )

    table_hdr_style = ParagraphStyle(
        'TableHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white
    )

    story = []

    # Title
    story.append(Paragraph("Blockchain Technology (CSE4444)", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("Assignment 2: Decentralized Application (DApp) Submission", sub_style))
    story.append(Spacer(1, 14))

    # Details Table
    details_data = [
        [Paragraph("<b>Student Name:</b>", body_style), Paragraph("<i>[Student Name]</i>", body_style)],
        [Paragraph("<b>Registration Number:</b>", body_style), Paragraph("<i>[Registration Number]</i>", body_style)],
        [Paragraph("<b>Assignment Title:</b>", body_style), Paragraph("Blockchain Technology (CSE4444) - Assignment 2", body_style)],
        [Paragraph("<b>Title of Developed DApp:</b>", body_style), Paragraph("Blockchain-Based Student Attendance Management System", body_style)],
        [Paragraph("<b>GitHub Repository Link:</b>", body_style), Paragraph("https://github.com/[YourUsername]/blockchain-student-attendance", body_style)]
    ]

    t_details = Table(details_data, colWidths=[150, 350])
    t_details.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,-1), colors.HexColor("#F1F5F9")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(t_details)
    story.append(Spacer(1, 12))

    # Section 1
    story.append(Paragraph("1. Problem Statement & System Overview", h1_style))
    story.append(Paragraph(
        "Traditional educational attendance tracking relies on manual paperwork or centralized databases vulnerable to "
        "retroactive modifications, proxy logging, and lack of transparency. The <b>Blockchain-Based Student Attendance Management System</b> "
        "is an original decentralized application built on Ethereum. Using smart contracts, attendance logs are secured immutably, "
        "allowing only the authorized instructor to record attendance while enabling students to verify their genuine records transparently.",
        body_style
    ))
    story.append(Spacer(1, 10))

    # Section 2
    story.append(Paragraph("2. Smart Contract Architecture", h1_style))
    story.append(Paragraph(
        "The smart contract is written in Solidity (v0.8.20) in <code>contracts/AttendanceSystem.sol</code>. Key features include:<br/>"
        "• <b>Access Control:</b> <code>onlyInstructor</code> modifier restricts registration and attendance logging strictly to the deployer.<br/>"
        "• <b>Validation:</b> Reverts on zero roll numbers, empty strings, duplicate registrations, and duplicate attendance dates.<br/>"
        "• <b>Events:</b> Emits <code>StudentRegistered</code>, <code>AttendanceMarked</code>, and <code>LectureConducted</code>.<br/>"
        "• <b>View Functions:</b> <code>getStudent</code>, <code>isAttendedOnDate</code>, <code>getAttendancePercentage</code>, and <code>getAllStudents</code>.",
        body_style
    ))
    story.append(Spacer(1, 10))

    # Section 3
    story.append(Paragraph("3. Frontend & MetaMask Integration", h1_style))
    story.append(Paragraph(
        "A responsive user interface was created with HTML5, CSS3, and JavaScript, integrated with Ethers.js v5.7.2. "
        "MetaMask is used for account authentication and transaction signing. The DApp dynamically detects network changes, "
        "verifies the instructor role, submits write transactions to the blockchain, and queries live ledger data without gas cost.",
        body_style
    ))
    story.append(Spacer(1, 10))

    # Section 4: Test Cases
    story.append(Paragraph("4. Comprehensive Testing & Results", h1_style))
    test_rows = [
        [
            Paragraph("TC", table_hdr_style),
            Paragraph("Category", table_hdr_style),
            Paragraph("Scenario / Revert Condition", table_hdr_style),
            Paragraph("Observed Result", table_hdr_style)
        ],
        [Paragraph("TC-01", table_cell_style), Paragraph("Valid", table_cell_style), Paragraph("Deploy with courseName parameter", table_cell_style), Paragraph("<b>PASS</b> - Deployed", table_cell_style)],
        [Paragraph("TC-02", table_cell_style), Paragraph("Valid", table_cell_style), Paragraph("Register student (101, 'Alice Johnson')", table_cell_style), Paragraph("<b>PASS</b> - Registered", table_cell_style)],
        [Paragraph("TC-03", table_cell_style), Paragraph("Valid", table_cell_style), Paragraph("Mark attendance (101, 20261017, Present)", table_cell_style), Paragraph("<b>PASS</b> - Recorded", table_cell_style)],
        [Paragraph("TC-04", table_cell_style), Paragraph("Valid", table_cell_style), Paragraph("Query student record and attendance %", table_cell_style), Paragraph("<b>PASS</b> - Displayed", table_cell_style)],
        [Paragraph("TC-05", table_cell_style), Paragraph("Invalid", table_cell_style), Paragraph("Register student with empty name ('')", table_cell_style), Paragraph("<b>PASS</b> - Reverted", table_cell_style)],
        [Paragraph("TC-06", table_cell_style), Paragraph("Invalid", table_cell_style), Paragraph("Register student with roll number 0", table_cell_style), Paragraph("<b>PASS</b> - Reverted", table_cell_style)],
        [Paragraph("TC-07", table_cell_style), Paragraph("Invalid", table_cell_style), Paragraph("Duplicate student registration (101)", table_cell_style), Paragraph("<b>PASS</b> - Reverted", table_cell_style)],
        [Paragraph("TC-08", table_cell_style), Paragraph("Invalid", table_cell_style), Paragraph("Attendance for unregistered student (999)", table_cell_style), Paragraph("<b>PASS</b> - Reverted", table_cell_style)],
        [Paragraph("TC-09", table_cell_style), Paragraph("Invalid", table_cell_style), Paragraph("Duplicate attendance on same date", table_cell_style), Paragraph("<b>PASS</b> - Reverted", table_cell_style)],
        [Paragraph("TC-10", table_cell_style), Paragraph("Unauthorized", table_cell_style), Paragraph("Non-instructor calls registerStudent", table_cell_style), Paragraph("<b>PASS</b> - Reverted", table_cell_style)],
        [Paragraph("TC-11", table_cell_style), Paragraph("Unauthorized", table_cell_style), Paragraph("Non-instructor calls markAttendance", table_cell_style), Paragraph("<b>PASS</b> - Reverted", table_cell_style)]
    ]

    t_test_pdf = Table(test_rows, colWidths=[40, 75, 260, 125])
    t_test_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1E3A8A")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t_test_pdf)
    story.append(Spacer(1, 10))

    # Section 5: Submission Checklist
    story.append(Paragraph("5. Submission Checklist & Repository Organization", h1_style))
    story.append(Paragraph(
        "• <code>contracts/AttendanceSystem.sol</code>: Solidity Smart Contract source code.<br/>"
        "• <code>frontend/</code>: index.html, style.css, app.js, and contract_abi.js.<br/>"
        "• <code>deployment/deployment_details.md</code>: Remix IDE, Ganache, and Sepolia deployment steps.<br/>"
        "• <code>test/test_contract.py & test_results.md</code>: Complete test suite and test results.<br/>"
        "• <code>screenshots/</code>: Compilation, deployment, MetaMask, and transaction verification evidence.<br/>"
        "• <code>README.md</code>: Project execution, setup guide, and contract interaction documentation.",
        body_style
    ))

    doc.build(story)
    print(f"PDF document saved to: {filepath}")

def main():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    docx_path = os.path.join(base_dir, "YourName_RegNum_BT_Assignment2.docx")
    pdf_path = os.path.join(base_dir, "YourName_RegNum_BT_Assignment2.pdf")

    generate_word_doc(docx_path)
    generate_pdf_doc(pdf_path)
    print("Both submission documents generated successfully!")

if __name__ == "__main__":
    main()
